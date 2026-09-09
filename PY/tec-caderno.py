#!/usr/bin/env python3
"""
Lê um export "Desempenho" do TecConcursos (.xlsx) e mostra, por disciplina:
o total, a árvore de tópicos-folha com os erros marcados, e uma sugestão de
`materia`/`bloco` casada com o vocabulário que o vault já usa.

Existe pra não precisar despejar 40+ linhas de planilha no contexto e ler
hierarquia a olho. Uso:

    python3 PY/tec-caderno.py ~/Downloads/<export>.xlsx

O que o script NÃO faz de propósito: não escreve nota nenhuma, não adivinha
data (o export não carrega data) e não decide matéria quando o casamento é
fraco — ele marca `??` e deixa a decisão para quem está conduzindo.
"""
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("falta openpyxl: python3 -m pip install openpyxl")
    sys.exit(1)

DIARIO = Path("Questoes/Diario")
MATERIAS = Path("MATERIAS")
STOP = {"de", "da", "do", "das", "dos", "e", "a", "o", "as", "os", "em", "com",
        "ao", "à", "às", "aos", "para", "por", "na", "no", "nas", "nos", "sem",
        "sobre", "entre"}


def tokens(s):
    s = unicodedata.normalize("NFKD", s.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"\(.*?\)", " ", s)
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return {w for w in s.split() if len(w) > 2 and w not in STOP}


def vocabulario():
    """materia -> bloco, a partir do que o Diario já usa (fonte de verdade)."""
    pares = Counter()
    for f in DIARIO.glob("*.md"):
        t = f.read_text(encoding="utf-8")
        m = re.search(r"^materia:[ \t]*(.*)$", t, re.M)
        b = re.search(r"^bloco:[ \t]*(.*)$", t, re.M)
        if m and m.group(1).strip():
            pares[(m.group(1).strip(), b.group(1).strip() if b else "")] += 1
    # para cada materia, o bloco mais frequente e não-vazio
    melhor = {}
    for (mat, blo), n in pares.most_common():
        if mat not in melhor or (not melhor[mat][0] and blo):
            melhor[mat] = (blo, n)
    return {mat: blo for mat, (blo, _) in melhor.items()}


def sugere(disciplina, vocab):
    tt = tokens(disciplina)
    if not tt:
        return None, 0.0
    alvo, score = None, 0.0
    for mat in vocab:
        mt = tokens(mat)
        if not mt:
            continue
        inter = len(tt & mt)
        s = max(inter / len(tt | mt), (inter / len(mt)) * 0.9)
        if s > score:
            alvo, score = mat, s
    return alvo, score


def carrega(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb["Desempenho"] if "Desempenho" in wb.sheetnames else wb[wb.sheetnames[0]]
    linhas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        hier = str(r[0]).strip() if r[0] is not None else ""
        linhas.append({
            "hier": hier,
            "nome": str(r[1]).strip(),
            "total": int(r[2] or 0),
            "acertos": int(r[4] or 0),
            "erros": int(r[6] or 0),
        })
    return linhas


def disciplinas(linhas):
    """Agrupa por disciplina (linha sem Hierarquia) e marca folhas."""
    grupos, atual = [], None
    for l in linhas:
        if not l["hier"]:
            atual = {"raiz": l, "itens": []}
            grupos.append(atual)
        elif atual is not None:
            atual["itens"].append(l)
    for g in grupos:
        hiers = [i["hier"] for i in g["itens"]]
        for i in g["itens"]:
            i["folha"] = not any(h.startswith(i["hier"] + ".") for h in hiers)
    return grupos


def main(caminho):
    linhas = carrega(caminho)
    grupos = disciplinas(linhas)
    vocab = vocabulario()
    total_geral = sum(g["raiz"]["total"] for g in grupos)

    print(f"{len(linhas)} linhas · {len(grupos)} disciplinas · {total_geral} questões no total\n")
    if total_geral > 120:
        print("⚠️  Volume alto para um dia só. O export do TEC também serve como")
        print("    acumulado de período/vida inteira — confirme o filtro antes de")
        print("    tratar isso como cadernos de um dia.\n")

    for g in grupos:
        raiz = g["raiz"]
        pct = round(100 * raiz["acertos"] / raiz["total"]) if raiz["total"] else 0
        mat, score = sugere(raiz["nome"], vocab)
        if score >= 0.5:
            casou = f"materia: {mat}   bloco: {vocab[mat] or '(VAZIO no Diario — preencher)'}"
        else:
            casou = f"?? sem casamento confiável (melhor palpite: {mat!r}, score {score:.2f}) — PERGUNTAR"
        print(f"=== {raiz['nome']} — {raiz['acertos']}/{raiz['total']} ({pct}%) ===")
        print(f"    {casou}")

        folhas = [i for i in g["itens"] if i["folha"]]
        soma_t = sum(i["total"] for i in folhas)
        soma_a = sum(i["acertos"] for i in folhas)
        print(f"    folhas (assuntos reais para o campo `assuntos:`):")
        for i in folhas:
            marca = "✗" if i["erros"] else " "
            print(f"      {marca} {i['acertos']}/{i['total']}  {i['nome']}")
        ok = "confere" if (soma_t, soma_a) == (raiz["total"], raiz["acertos"]) else "NÃO CONFERE"
        print(f"    soma das folhas: {soma_a}/{soma_t}  [{ok} com o total da disciplina]")

        com_erro = [i for i in folhas if i["erros"]]
        if com_erro:
            print(f"    erros para a seção `## Erros a revisar`:")
            for i in com_erro:
                print(f"      - {i['nome']} — {i['acertos']}/{i['total']}")
        else:
            print("    sem erros → deixe `erro_tipo` VAZIO (o painel Diagnóstico de erro")
            print("      filtra por presença do campo; preenchê-lo aqui lista a matéria")
            print("      como lacuna de conhecimento sem nenhum erro a contabilizar)")
        print()

    print(f"TOTAL A CONCILIAR: {sum(g['raiz']['acertos'] for g in grupos)}/{total_geral}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("uso: python3 PY/tec-caderno.py <export.xlsx>")
        sys.exit(1)
    main(Path(sys.argv[1]).expanduser())
